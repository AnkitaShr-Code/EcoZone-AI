import sys, zipfile, re
from pptx import Presentation
p = Presentation(sys.argv[1]); SW, SH = p.slide_width, p.slide_height
bad_phrases = ["Please download", "How different is it", "How will it be able", "USP of the proposed", "Kindly refrain"]
problems = 0
for i, s in enumerate(p.slides, 1):
    texts = []
    for sh in s.shapes:
        l, t, w, h = sh.left, sh.top, sh.width, sh.height
        # template bars sit at -0.03 in and 5.55..5.66 in: tolerate up to 0.04 in.
        tol = 914400 * 0.04
        if l < -tol or t < -tol or l + w > SW + tol or t + h > SH + tol:
            print(f"  slide {i}: {sh.name} out of bounds {l/914400:.2f},{t/914400:.2f} {w/914400:.2f}x{h/914400:.2f}"); problems += 1
        if sh.is_placeholder and sh.has_text_frame and not sh.text_frame.text.strip():
            print(f"  slide {i}: empty placeholder {sh.name}"); problems += 1
        if sh.has_text_frame and sh.text_frame.text.strip():
            texts.append(sh.text_frame.text.strip().replace("\n", " | "))
        if sh.shape_type == 13:
            _ = sh.image.blob  # raises if the relationship is broken
    joined = " ".join(texts)
    for b in bad_phrases:
        if b in joined:
            print(f"  slide {i}: leftover template text '{b}'"); problems += 1
    pics = sum(1 for sh in s.shapes if sh.shape_type == 13)
    print(f"{i:2d} pics={pics} text: {joined[:150]}")
z = zipfile.ZipFile(sys.argv[1]); print("zip ok:", z.testzip() is None, "slides in zip:", len([n for n in z.namelist() if re.match(r'ppt/slides/slide\d+\.xml$', n)]))
print("problems:", problems)
