"""Fill the AI Builder Cup template with the EcoZone AI deck.

Reads the template from the repo root (never written), the 2x renders in deck-work/hires/, and writes
deck-work/NeuralNodes-AI-Builder-Cup-Submission.pptx. Run: uv run --with python-pptx --with pillow python deck-work/build_pptx.py
"""

import copy
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Emu, Inches, Pt
from pptx.dml.color import RGBColor

ROOT = Path("/Users/abhi/work/EcoZone-AI")
W = ROOT / "deck-work"
TEMPLATE = ROOT / "Submission Template _ AI Builder Cup.pptx"
OUT = W / "NeuralNodes-AI-Builder-Cup-Submission.pptx"
HIRES = W / "hires"
CROPS = W / "pptx-img"
CROPS.mkdir(exist_ok=True)

GITHUB = "https://github.com/AnkitaShr-Code/EcoZone-AI"
LIVE = "https://ecozone-451227459430.asia-south1.run.app/"
FONT = "Google Sans Flex"

prs = Presentation(str(TEMPLATE))
SW, SH = prs.slide_width, prs.slide_height
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


def crop(slide_id: str, top: int, bottom: int) -> Path:
    """Crop a 2x render between 1x pixel rows top..bottom and save as high-quality JPEG."""
    im = Image.open(HIRES / f"{slide_id}.png").convert("RGB")
    out = CROPS / f"{slide_id}.jpg"
    im.crop((0, top * 2, im.width, bottom * 2)).save(out, quality=92, optimize=True)
    return out


def duplicate(src):
    """Append a copy of a slide (shapes and image relationships) and return it."""
    new = prs.slides.add_slide(src.slide_layout)
    for shp in list(new.shapes):
        shp._element.getparent().remove(shp._element)
    rid_map = {}
    for rid, rel in src.part.rels.items():
        if rel.reltype.endswith("/image"):
            rid_map[rid] = new.part.relate_to(rel._target, rel.reltype)
    for shp in src.shapes:
        el = copy.deepcopy(shp._element)
        for node in el.iter():
            for attr in (f"{R}embed", f"{R}link", f"{R}id"):
                if node.get(attr) in rid_map:
                    node.set(attr, rid_map[node.get(attr)])
        new.shapes._spTree.append(el)
    return new


def move(slide, index: int) -> None:
    lst = prs.slides._sldIdLst
    ids = list(lst)
    el = next(i for i in ids if prs.slides.part.related_part(i.get(f"{R}id")) is slide.part)
    lst.remove(el)
    lst.insert(index, el)


def delete(slide) -> None:
    lst = prs.slides._sldIdLst
    el = next(i for i in lst if prs.slides.part.related_part(i.get(f"{R}id")) is slide.part)
    prs.part.drop_rel(el.get(f"{R}id"))
    lst.remove(el)


def title_box(slide):
    """The template's title text box: the text box nearest y = 0.73 in."""
    boxes = [s for s in slide.shapes if s.has_text_frame and s.shape_type == 17]
    return min(boxes, key=lambda s: abs(s.top - Inches(0.73)))


def set_paragraph(par, text: str, url: str | None = None, size: int | None = None) -> None:
    runs = par.runs
    if not runs:
        run = par.add_run()
    else:
        run = runs[0]
        for extra in runs[1:]:
            extra._r.getparent().remove(extra._r)
    run.text = text
    if size:
        run.font.size = Pt(size)
    if url:
        run.hyperlink.address = url


def set_title(slide, text: str) -> None:
    tf = title_box(slide).text_frame
    set_paragraph(tf.paragraphs[0], text)
    for par in tf.paragraphs[1:]:
        par._p.getparent().remove(par._p)


def add_picture(slide, path: Path, top_in: float, max_h_in: float) -> None:
    with Image.open(path) as im:
        aspect = im.width / im.height
    h = Inches(max_h_in)
    w = int(h * aspect)
    if w > SW:
        w = SW
        h = int(w / aspect)
    left = int((SW - w) / 2)
    slide.shapes.add_picture(str(path), left, Inches(top_in), w, h)


def add_text(slide, left, top, width, height, lines, color=RGBColor(0xFF, 0xFF, 0xFF)):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    for i, (text, size, bold) in enumerate(lines):
        par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = par.add_run()
        run.text = text
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.name = FONT
        run.font.color.rgb = color
    return box


slides = list(prs.slides)
(s_team, s_note, s_brief, s_opp, s_feat, s_proc, s_wire, s_arch, s_tech, s_cost, s_snap, s_perf, s_future,
 s_links, s_close, s_thanks) = slides

# Continuation pages for snapshots (2) and performance (3), placed right after their template slide.
s_snap2 = duplicate(s_snap)
s_perf2 = duplicate(s_perf)
s_perf3 = duplicate(s_perf)
delete(s_note)
order = [s_team, s_brief, s_opp, s_feat, s_proc, s_wire, s_arch, s_tech, s_cost, s_snap, s_snap2, s_perf, s_perf2,
         s_perf3, s_future, s_links, s_close, s_thanks]
for i, s in enumerate(order):
    move(s, i)

# 1. Team Details: the template's own fields only (team name, team leader, problem statement).
tf = next(s for s in s_team.shapes if s.has_text_frame).text_frame
pars = tf.paragraphs
set_paragraph(pars[2], "Team name: NeuralNodes")
set_paragraph(pars[3], "Team leader name: Ankita Sharma")
set_paragraph(pars[4], "Problem Statement: Cities have no quick way to decide where trees, shade, depaving and cool "
                       "roofs would cool the people who wait outdoors, or to stop new buildings breaking the breeze "
                       "and green that cool a street. Theme: Sustainability & Social Impact.", size=16)

# Content slides: template title kept as text, our slide body (subtitle, charts, screenshots) as one picture.
content = [
    (s_brief, "brief", "Brief about the idea"),
    (s_opp, "opportunities", "Opportunities"),
    (s_feat, "features", "List of features offered by the solution"),
    (s_proc, "process", "Process flow diagram or Use-case diagram"),
    (s_wire, "wireframe", "Wireframes/Mock diagrams of the proposed solution (optional)"),
    (s_arch, "architecture", "Architecture diagram of the proposed solution"),
    (s_tech, "tech", "Technologies to be used in the solution"),
    (s_cost, "cost", "Estimated implementation cost (optional)"),
    (s_snap, "snaps-repair", "Snapshots of the prototype · Repair"),
    (s_snap2, "snaps-protect", "Snapshots of the prototype · Protect"),
    (s_perf, "perf-evals", "Prototype Performance report/Benchmarking · 1 of 3"),
    (s_perf2, "perf-gemini", "Prototype Performance report/Benchmarking · 2 of 3"),
    (s_perf3, "perf-heat", "Prototype Performance report/Benchmarking · 3 of 3"),
    (s_future, "future", "Additional Details/Future Development (if any)"),
]
for slide, sid, title in content:
    set_title(slide, title)
    add_picture(slide, crop(sid, 188, 1052), top_in=1.08, max_h_in=4.45)

# 14. Links: native, editable text with clickable links (repo, demo video on Google Drive, live app).
tf = title_box(s_links).text_frame
pars = tf.paragraphs
set_paragraph(pars[0], "Provide links to your:")
set_paragraph(pars[2], f"GitHub Public Repository: {GITHUB}", url=GITHUB)
DEMO = "https://drive.google.com/file/d/18SEjtILQqo1EFyp77QhDfIrvQbLatgSB/view?usp=sharing"
set_paragraph(pars[3], "Demo Video Link (3 Minutes): Watch the demo video (Google Drive)", url=DEMO)
set_paragraph(pars[4], f"Final Product Link: {LIVE}", url=LIVE)
add_picture(s_links, crop("links", 705, 935), top_in=3.2, max_h_in=1.25)

# 15. Closing statement (content of our closing slide, below the template header strip).
add_picture(s_close, crop("close", 112, 1052), top_in=0.58, max_h_in=4.9)

# 16. Thank you: template artwork plus the team leader and the two links.
add_text(s_thanks, 6.75, 4.1, 3.1, 1.2, [
    ("Ankita Sharma", 16, True),
    ("Team leader, NeuralNodes", 12, False),
    ("github.com/AnkitaShr-Code/EcoZone-AI", 10, False),
    ("ecozone-451227459430.asia-south1.run.app", 10, False),
])

# Keep every shape inside the 10 x 5.625 in slide: the template's title boxes run 0.07 in past the right edge.
for slide in prs.slides:
    for shp in slide.shapes:
        if shp.has_text_frame and shp.left + shp.width > SW:
            shp.width = SW - shp.left - Inches(0.05)

prs.save(str(OUT))
print("wrote", OUT, OUT.stat().st_size // 1024, "KB,", len(prs.slides), "slides")
